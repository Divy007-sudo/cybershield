"""Run with EXCEL_FILE and SOURCE_CSV pointing to the download and original CSV."""
import csv
import os
import unittest
from openpyxl import load_workbook

COLUMNS = ['id', 'dur', 'proto', 'service', 'state', 'spkts', 'dpkts', 'sbytes', 'dbytes', 'rate', 'sttl', 'attack_cat', 'label']
TEXT = {'proto', 'service', 'state', 'attack_cat'}


class ExcelDownloadTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        wb = load_workbook(os.environ['EXCEL_FILE'], read_only=True, data_only=True)
        cls.sheets = wb.sheetnames
        cls.values = list(wb.active.values)
        wb.close()

    def test_exactly_175341_records_plus_header(self):
        self.assertEqual(len(self.values), 175342)
        self.assertEqual(len(self.sheets), 1)

    def test_exact_13_columns_in_requested_order(self):
        self.assertEqual(list(self.values[0]), COLUMNS)
        self.assertTrue(all(len(row) == 13 for row in self.values))

    def test_every_record_matches_full_actual_source(self):
        with open(os.environ['SOURCE_CSV'], encoding='utf-8-sig', newline='') as f:
            source = list(csv.DictReader(f))
        self.assertEqual(len(source), 175341)
        for index, (original, exported) in enumerate(zip(source, self.values[1:]), 1):
            expected = tuple(original[c].strip() if c in TEXT else float(original[c]) for c in COLUMNS)
            self.assertEqual(exported, expected, f'Source row {index} differs')


if __name__ == '__main__':
    unittest.main()